#!/usr/bin/env python3
"""Arma el trabajo de grado en la plantilla NTC 1486 (pregrado UCC, 28 jul 2026)."""
from __future__ import annotations

import json
import re
import shutil
import struct
from pathlib import Path

import pymupdf
from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Emu, Inches, Pt, RGBColor, Twips

TPL = Path(
    "/home/ubuntu/.cursor/projects/workspace/uploads/"
    "Plantilla_Documento_final_28_de_julio_de_2026_Pregrado__2__1f55.docx"
)
PDF = Path(
    "/home/ubuntu/.cursor/projects/workspace/uploads/"
    "Trabajo_de_Grado_La_Manuelita_actualizado__4__1050.pdf"
)
SRC = Path("/tmp/tesis-v4.txt")
ASSET = Path("/tmp/tesis-ntc-assets")
OUT = Path("/workspace/Trabajo_de_Grado_La_Manuelita_NTC1486.docx")
ART = Path("/opt/cursor/artifacts/Trabajo_de_Grado_La_Manuelita_NTC1486.docx")
DOCS = Path("/workspace/docs/tesis/Trabajo_de_Grado_La_Manuelita_NTC1486.docx")

TITLE = (
    "DISEÑO E IMPLEMENTACIÓN DE UN PROTOTIPO DE SISTEMA DE INFORMACIÓN "
    "OFFLINE-FIRST PARA OPTIMIZAR LA OPERACIÓN DE UNA LAVANDERÍA EN BOGOTÁ D.C."
)

FIG_XREF = {
    1: 628, 2: 678, 3: 1519, 4: 1523, 5: 1564, 6: 1659, 7: 1776,
    8: 1782, 9: 1791, 10: 1795, 11: 1803, 12: 1808, 13: 1815, 14: 1819,
    15: 1828, 16: 1837, 17: 1845,
}
ANEXO_XREF = {
    "A1": 1999, "A2": 2002, "B": 2012, "C": 2019, "D": 2026, "F": 2037,
}

TAB_CAPTION = {
    1: "Matriz comparativa del estado del arte",
    2: "Indicadores clave de desempeño (KPI)",
    3: "Matriz de requerimientos funcionales y no funcionales",
    4: "Tarifario de servicios — Lavaseco La Manuelita",
    5: "Línea base medida — reemplaza la línea base declarada",
    6: "Modelo de datos del prototipo",
    7: "Estado de los módulos del prototipo",
}
TAB_CONT_CAPTION = "Módulo asociado y criterio de validación por requerimiento"

FIG_CAPTION = {
    1: "Flujo del proceso operativo actual — Lavaseco La Manuelita. Elaboración propia.",
    2: "Diagrama de casos de uso — Lavaseco La Manuelita. Elaboración propia.",
    3: "Causas de confusión de pedidos — Lavaseco La Manuelita. Elaboración propia.",
    4: "Línea base declarada vs. medida — Lavaseco La Manuelita. Elaboración propia.",
    5: "Arquitectura de despliegue local-first — Lavaseco La Manuelita. Elaboración propia.",
    6: (
        "Diagrama entidad-relación del prototipo — Lavaseco La Manuelita. Elaboración propia. "
        "Nota: Tarifario no tiene llave foránea hacia Prenda; el cálculo del valor se resuelve "
        "por coincidencia de texto en el cliente (numeral 3.2)."
    ),
    7: "Diagrama de estados del pedido — Lavaseco La Manuelita. Elaboración propia.",
    8: "Acceso al sistema — portal de empleado y panel de gerente (RNF04). Captura del prototipo.",
    9: "Pedido rápido — paso 1: datos del cliente (RF01). Captura del prototipo.",
    10: "Panel financiero — ingresos diarios y pagos por método (RF09 y reportería adicional). Captura del prototipo.",
    11: "Historial de cierres de caja, con entradas y salidas del día (RF09). Captura del prototipo.",
    12: "Pedidos con más de tres meses en piso (RF11). Captura del prototipo.",
    13: "Inventario en piso, con filtros y acceso a escaneo de código de barras (RF06, RF03). Captura del prototipo.",
    14: "Detalle de pedido — prendas, abono y cambio de estado (RF03, RF04). Captura del prototipo.",
    15: "Recibo impreso con código de barras — soporte de trazabilidad de la marquilla. Captura del prototipo.",
    16: "Panel remoto de la gerente, con conexión disponible.",
    17: "Panel remoto de la gerente, sin conexión (modo avión).",
}

ANEXO_CAPTION = {
    "A": "Formato de recibo físico utilizado actualmente por Lavaseco La Manuelita (dos ejemplos, con nombre y teléfono anonimizados).",
    "B": "Página del cuaderno de salidas — registro de recibos tachados al retirarse el pedido del piso (sin datos personales).",
    "C": "Página del cuaderno de entradas — registro consolidado de pedidos recibidos (recibos No. 33791–33819), con nombres de clientes anonimizados.",
    "D": "Página del cuaderno de entradas — registro consolidado de pedidos recibidos (recibos No. 33767–33790), con nombres de clientes anonimizados.",
    "E": "Guion de la entrevista semiestructurada aplicada a la gerente. Documento completo adjunto por separado (Anexo_E_Guion_Entrevista_Semiestructurada.docx).",
    "F": "Carta de autorización del negocio para el uso de su información en el trabajo de grado, firmada por la gerente Helena del Carmen Contreras el 10 de septiembre de 2026.",
}

PALABRAS_CLAVE = (
    "almacenamiento de datos, computación en la nube, pequeña empresa, "
    "procesamiento de datos, sincronización de datos, sistema de información, software de aplicación."
)
KEYWORDS = (
    "application software, cloud computing, data processing, data storage, "
    "data synchronization, information system, small business."
)

KEEP_HYPHEN = {
    "offline", "local", "black", "post", "web", "real", "non", "multi", "micro",
    "end", "front", "data", "cloud", "open", "self", "read", "write", "e",
}


def png_size(path: Path) -> tuple[int, int]:
    with path.open("rb") as f:
        sig = f.read(8)
        if sig != b"\x89PNG\r\n\x1a\n":
            return (1600, 900)
        length, ctype = struct.unpack(">I4s", f.read(8))
        if ctype != b"IHDR":
            return (1600, 900)
        w, h = struct.unpack(">II", f.read(8))
        return w, h


def save_xref(pdf: pymupdf.Document, xref: int, path: Path) -> None:
    pix = pymupdf.Pixmap(pdf, xref)
    if pix.n - pix.alpha > 3:
        pix = pymupdf.Pixmap(pymupdf.csRGB, pix)
    path.parent.mkdir(parents=True, exist_ok=True)
    pix.save(str(path))


def cell_text(s: str | None) -> str:
    if not s:
        return ""
    t = re.sub(r"[ \t]*\n[ \t]*", " ", s)
    t = re.sub(r"\s+", " ", t).strip()
    t = t.replace("prisma.tarifario", "el catálogo de tarifas almacenado en la base de datos")
    t = t.replace(" (lib/caja.ts)", "")
    return t


def extract_table(page, index=0) -> list[list[str]]:
    tabs = page.find_tables()
    if not tabs or index >= len(tabs.tables):
        return []
    data = tabs.tables[index].extract()
    return [[cell_text(c) for c in row] for row in data]


def merge_tables(chunks: list[list[list[str]]]) -> list[list[str]]:
    rows: list[list[str]] = []
    header = None
    for i, data in enumerate(chunks):
        if not data:
            continue
        if i == 0:
            rows.extend(data)
            header = data[0][0] if data else None
        else:
            if data[0][0] == header:
                rows.extend(data[1:])
            else:
                rows.extend(data)
    return rows


def extract_assets() -> dict:
    ASSET.mkdir(parents=True, exist_ok=True)
    pdf = pymupdf.open(str(PDF))
    figs = {}
    for n, xref in FIG_XREF.items():
        p = ASSET / f"figura_{n:02d}.png"
        save_xref(pdf, xref, p)
        figs[n] = p
    anex = {}
    for k, xref in ANEXO_XREF.items():
        p = ASSET / f"anexo_{k}.png"
        save_xref(pdf, xref, p)
        anex[k] = p

    t1 = merge_tables([
        extract_table(pdf[20]),
        extract_table(pdf[21]),
        extract_table(pdf[22]),
    ])
    t2 = extract_table(pdf[23])
    t3 = merge_tables([
        extract_table(pdf[28]),
        extract_table(pdf[29]),
        extract_table(pdf[30]),
        extract_table(pdf[31], 0),
    ])
    t3c = merge_tables([
        extract_table(pdf[31], 1),
        extract_table(pdf[32]),
        extract_table(pdf[33], 0),
    ])
    t4 = merge_tables([
        extract_table(pdf[33], 1),
        extract_table(pdf[34]),
    ])
    t5 = extract_table(pdf[35])
    t6 = merge_tables([
        extract_table(pdf[41]),
        extract_table(pdf[42]),
    ])
    t7 = merge_tables([
        extract_table(pdf[43]),
        extract_table(pdf[44]),
    ])
    tables = {"1": t1, "2": t2, "3": t3, "3c": t3c, "4": t4, "5": t5, "6": t6, "7": t7}
    (ASSET / "tables.json").write_text(json.dumps(tables, ensure_ascii=False), encoding="utf-8")

    if not SRC.exists():
        parts = []
        for i, page in enumerate(pdf):
            parts.append(f"\n===== PAGE {i+1}/{pdf.page_count} =====\n")
            parts.append(page.get_text("text"))
        SRC.write_text("".join(parts), encoding="utf-8")
    pdf.close()
    return {"figs": figs, "anex": anex, "tables": tables}


def add_field(paragraph, instr: str, placeholder: str = "") -> None:
    r1 = paragraph.add_run()
    b = OxmlElement("w:fldChar")
    b.set(qn("w:fldCharType"), "begin")
    r1._r.append(b)
    r2 = paragraph.add_run()
    it = OxmlElement("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = " " + instr.strip() + " "
    r2._r.append(it)
    r3 = paragraph.add_run()
    s = OxmlElement("w:fldChar")
    s.set(qn("w:fldCharType"), "separate")
    r3._r.append(s)
    if placeholder:
        paragraph.add_run(placeholder)
    r4 = paragraph.add_run()
    e = OxmlElement("w:fldChar")
    e.set(qn("w:fldCharType"), "end")
    r4._r.append(e)


def set_run_arial(run, size_pt=12, bold=False, italic=False):
    run.font.name = "Arial"
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:ascii"), "Arial")
    rFonts.set(qn("w:hAnsi"), "Arial")
    rFonts.set(qn("w:eastAsia"), "Arial")
    rFonts.set(qn("w:cs"), "Arial")
    run.font.size = Pt(size_pt)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor(0, 0, 0)


def add_centered(doc, text, size=12, bold=False, space_before=0, space_after=0):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.first_line_indent = Cm(0)
    run = p.add_run(text)
    set_run_arial(run, size, bold)
    return p


def add_body(doc, text, indent=True):
    p = doc.add_paragraph(style="Normal")
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.first_line_indent = Cm(1.25) if indent else Cm(0)
    run = p.add_run(text)
    set_run_arial(run, 12)
    return p


def add_front_title(doc, text):
    """Título preliminar con apariencia de Heading 1, sin entrar al esquema del TOC."""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(12)
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.page_break_before = text != "CONTENIDO"
    run = p.add_run(text.upper())
    set_run_arial(run, 12, bold=True)
    return p


def add_heading_styled(doc, text, level: int, page_break: bool | None = None):
    style = {1: "Heading 1", 2: "Heading 2", 3: "Heading 3"}[level]
    p = doc.add_paragraph(text, style=style)
    p.paragraph_format.space_before = Pt(18 if level == 1 else 12)
    p.paragraph_format.space_after = Pt(10)
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.line_spacing = 1.15
    if level == 1:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if page_break is None:
            page_break = text in {
                "GLOSARIO", "RESUMEN", "ABSTRACT",
                "INTRODUCCIÓN", "1. GENERALIDADES",
                "2. CARACTERIZACIÓN DEL PROCESO OPERATIVO DE LAVASECO LA MANUELITA",
                "3. DESARROLLO DEL SEGUNDO OBJETIVO",
                "4. DESARROLLO DEL TERCER OBJETIVO",
                "5. CONCLUSIONES", "6. RECOMENDACIONES",
                "BIBLIOGRAFÍA", "ANEXOS",
            }
        p.paragraph_format.page_break_before = bool(page_break)
    else:
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for run in p.runs:
        set_run_arial(run, 12, bold=True)
    return p


SEQ_N = {"Figura": 0, "Tabla": 0, "Anexo": 0}


def add_caption(doc, kind: str, title: str, continuation: bool = False):
    p = doc.add_paragraph(style="Caption")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.15
    r0 = p.add_run(f"{kind} ")
    set_run_arial(r0, 10, italic=True)
    if kind == "Anexo":
        SEQ_N[kind] += 1
        letter = chr(ord("A") + SEQ_N[kind] - 1)
        add_field(p, "SEQ Anexo \\* ALPHABETIC", letter)
        r1 = p.add_run(f". {title}")
    elif continuation:
        add_field(p, f"SEQ {kind} \\c", str(SEQ_N[kind]))
        r1 = p.add_run(f" (continuación). {title}")
    else:
        SEQ_N[kind] += 1
        add_field(p, f"SEQ {kind} \\* ARABIC", str(SEQ_N[kind]))
        r1 = p.add_run(f". {title}")
    set_run_arial(r1, 10, italic=True)
    return p


def add_image(doc, path: Path, max_w_cm=15.5, max_h_cm=20.0):
    if not path or not path.exists():
        return
    w_px, h_px = png_size(path)
    if w_px <= 0 or h_px <= 0:
        return
    max_w = Cm(max_w_cm)
    max_h = Cm(max_h_cm)
    ratio = h_px / w_px
    width = max_w
    height = int(width * ratio)
    if height > max_h:
        height = max_h
        width = int(height / ratio)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run()
    run.add_picture(str(path), width=width, height=height)


def page_break(doc):
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Cm(0)
    p.add_run().add_break(WD_BREAK.PAGE)


def shade_cell(cell, color: str):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color)
    tcPr.append(shd)


def set_cell_border(cell):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "000000")
        tcBorders.append(el)
    tcPr.append(tcBorders)


def set_repeat_header(row):
    tr = row._tr
    trPr = tr.get_or_add_trPr()
    tblHeader = OxmlElement("w:tblHeader")
    tblHeader.set(qn("w:val"), "true")
    trPr.append(tblHeader)


def add_word_table(doc, data: list[list[str]]):
    if not data:
        return
    cols = max(len(r) for r in data)
    table = doc.add_table(rows=len(data), cols=cols)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    for i, row in enumerate(data):
        for j in range(cols):
            cell = table.cell(i, j)
            cell.text = ""
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            txt = row[j] if j < len(row) else ""
            run = p.add_run(txt)
            set_run_arial(run, 9, bold=(i == 0))
            set_cell_border(cell)
            if i == 0:
                shade_cell(cell, "D9E2F3")
        if i == 0:
            set_repeat_header(table.rows[0])
    return table


def clean_source(raw: str) -> str:
    raw = re.sub(r"\n===== PAGE \d+/\d+ =====\n", "\n", raw)
    raw = re.sub(r"\n\d{1,3}\n", "\n", raw)
    raw = raw.replace("\u00ad", "")
    raw = raw.replace("prisma.tarifario", "el catálogo de tarifas almacenado en la base de datos")
    raw = raw.replace(" (lib/caja.ts)", "")
    raw = re.sub(r"/gerente\?\s*fecha=", "/gerente?fecha=", raw)
    raw = re.sub(r"\n{3,}", "\n\n", raw)
    return raw


def is_h1(line: str) -> bool:
    s = line.strip()
    if s in {
        "INTRODUCCIÓN", "1. GENERALIDADES",
        "3. DESARROLLO DEL SEGUNDO OBJETIVO",
        "4. DESARROLLO DEL TERCER OBJETIVO",
        "5. CONCLUSIONES", "6. RECOMENDACIONES",
        "BIBLIOGRAFÍA", "ANEXOS", "GLOSARIO", "RESUMEN", "ABSTRACT",
        "CONTENIDO", "LISTA DE TABLAS", "LISTA DE FIGURAS", "LISTA DE ANEXOS",
    }:
        return True
    if s.startswith("2. CARACTERIZACIÓN"):
        return True
    return False


def is_h2(line: str) -> bool:
    s = line.strip()
    return bool(re.match(r"^[1-3]\.[1-9] [A-ZÁÉÍÓÚÑa-záéíóúñ]", s))


def is_h3(line: str) -> bool:
    s = line.strip()
    return bool(re.match(r"^1\.[45]\.\d+ [A-ZÁÉÍÓÚÑ]", s))


def is_fig(line: str):
    return re.match(r"^Figura (\d+)\.\s*(.*)$", line.strip())


def is_tab(line: str):
    return re.match(r"^Tabla (\d+)(?: \(continuación\))?\.\s*(.*)$", line.strip())


def is_anexo(line: str):
    return re.match(r"^Anexo ([A-F])\.\s*(.*)$", line.strip())


def join_lines(lines: list[str]) -> str:
    buf = ""
    for ln in lines:
        s = ln.strip()
        if not s:
            continue
        if buf.endswith("-"):
            left = buf[:-1]
            last = left.split()[-1] if left.split() else left
            if last.lower().rstrip(".,;:") in KEEP_HYPHEN:
                buf = buf + s
            else:
                buf = left + s
        elif buf:
            buf += " " + s
        else:
            buf = s
    return buf


def _is_structure(s: str) -> bool:
    return bool(
        is_h1(s) or is_h2(s) or is_h3(s) or is_fig(s) or is_tab(s) or is_anexo(s)
        or s.startswith("2. CARACTERIZACIÓN")
        or s.startswith("●") or s.startswith("•")
        or s.startswith("Objetivo asociado")
        or s.startswith("PALABRAS CLAVE")
        or s.startswith("KEYWORDS")
    )


def _is_narrative_start(s: str) -> bool:
    if len(s) < 50:
        return False
    return s.startswith((
        "El ", "La ", "Los ", "Las ", "En ", "Un ", "Una ", "Este ", "Esta ",
        "Estos ", "Estas ", "Como ", "A partir", "Frente ", "De acuerdo",
        "Durante ", "Siguiendo ", "Se ", "Con ", "Para ", "Al ", "Adicionalmente",
        "Asimismo", "Además", "Por ", "Desde ", "Nota", "La Figura", "Las Figura",
        "El trabajo", "La caracterización", "La arquitectura", "El prototipo",
        "El desarrollo", "El primer", "El segundo", "El tercer", "El módulo",
        "El historial", "El autocompletado", "Insumo directo",
    ))


def reflow_paragraphs(chunk: str, biblio: bool = False) -> list[str]:
    lines = [ln.rstrip() for ln in chunk.splitlines()]
    paras: list[str] = []
    buf: list[str] = []

    def flush():
        nonlocal buf
        if buf:
            t = join_lines(buf).strip()
            if t:
                paras.append(t)
            buf = []

    def skip_table_body(start: int) -> int:
        i = start
        while i < len(lines):
            nxt = lines[i].strip()
            if not nxt:
                i += 1
                continue
            if _is_structure(nxt) or _is_narrative_start(nxt):
                return i
            i += 1
        return i

    i = 0
    while i < len(lines):
        s = lines[i].strip()
        if not s:
            flush()
            i += 1
            continue
        if s.startswith("(En ") or s.startswith("(Se "):
            flush()
            j = i + 1
            block = [s]
            while j < len(lines):
                nxt = lines[j].strip()
                if not nxt or _is_structure(nxt) or _is_narrative_start(nxt):
                    break
                block.append(nxt)
                j += 1
            paras.append(join_lines(block))
            i = j
            continue
        if s.startswith("2. CARACTERIZACIÓN"):
            flush()
            title = "2. CARACTERIZACIÓN DEL PROCESO OPERATIVO DE LAVASECO LA MANUELITA"
            if i + 1 < len(lines) and "LA MANUELITA" in lines[i + 1]:
                i += 1
            paras.append(title)
            i += 1
            continue
        if is_tab(s):
            flush()
            paras.append(s)
            i += 1
            while i < len(lines):
                nxt = lines[i].strip()
                if not nxt:
                    i += 1
                    continue
                if nxt.startswith("Insumo directo") or nxt.startswith("Cuando llega"):
                    note = [nxt]
                    i += 1
                    while i < len(lines):
                        n2 = lines[i].strip()
                        if not n2 or _is_structure(n2) or n2 in {
                            "Categoría", "Ítem", "Precio (COP)", "Indicador", "Código",
                            "Autor / año", "Objetivo específico", "Entidad", "Módulo",
                        }:
                            break
                        note.append(n2)
                        i += 1
                    paras.append(join_lines(note))
                    continue
                if _is_structure(nxt) or _is_narrative_start(nxt):
                    break
                i += 1
            continue
        if is_fig(s):
            flush()
            cap = s
            i += 1
            while i < len(lines):
                nxt = lines[i].strip()
                if not nxt:
                    break
                if _is_structure(nxt) or _is_narrative_start(nxt):
                    break
                if nxt.startswith("Nota") or cap.rstrip().endswith("Nota:"):
                    cap += " " + nxt
                    i += 1
                    continue
                break
            paras.append(cap)
            continue
        if is_h1(s) or is_h2(s) or is_h3(s):
            flush()
            paras.append(s)
            i += 1
            continue
        if is_anexo(s):
            flush()
            cap = s
            i += 1
            while i < len(lines):
                nxt = lines[i].strip()
                if not nxt or is_anexo(nxt) or is_h1(nxt):
                    break
                cap += " " + nxt
                i += 1
            paras.append(re.sub(r"\s+", " ", cap).strip())
            continue
        if s.startswith("●") or s.startswith("•"):
            flush()
            j = i + 1
            bullet = [s]
            while j < len(lines):
                nxt = lines[j].strip()
                if not nxt or nxt.startswith("●") or nxt.startswith("•") or _is_structure(nxt):
                    break
                bullet.append(nxt)
                j += 1
            paras.append(join_lines(bullet))
            i = j
            continue
        if biblio and re.match(r"^\[\d+\]", s):
            flush()
            j = i + 1
            block = [s]
            while j < len(lines):
                nxt = lines[j].strip()
                if not nxt:
                    j += 1
                    continue
                if re.match(r"^\[\d+\]", nxt) or is_h1(nxt) or is_anexo(nxt):
                    break
                block.append(nxt)
                j += 1
            paras.append(re.sub(r"\s+", " ", " ".join(block)).strip())
            i = j
            continue
        buf.append(s)
        i += 1
    flush()
    return paras


def parse_glossary(chunk: str) -> list[tuple[str, str]]:
    entries: list[tuple[str, str]] = []
    current_term = None
    buf: list[str] = []
    term_re = re.compile(r"^([A-ZÁÉÍÓÚÑ][A-ZÁÉÍÓÚÑ0-9 /().,-]+):\s*(.*)$")

    def flush():
        nonlocal current_term, buf
        if current_term and buf:
            d = join_lines(buf).strip()
            if d:
                d = d[0].lower() + d[1:]
            entries.append((current_term, d))
        current_term = None
        buf = []

    for ln in chunk.splitlines():
        s = ln.strip()
        if not s:
            continue
        m = term_re.match(s)
        if m:
            flush()
            current_term = m.group(1).strip()
            rest = m.group(2).strip()
            buf = [rest] if rest else []
        elif current_term:
            buf.append(s)
    flush()
    return entries


def extract_between(text: str, start: str, end: str | None) -> str:
    i = text.find(start)
    if i < 0:
        return ""
    i2 = text.find(end, i + len(start)) if end else len(text)
    if i2 < 0:
        i2 = len(text)
    return text[i + len(start): i2].strip()


def set_margins(section, left=3.0, right=2.0, top=3.0, bottom=2.0):
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.left_margin = Cm(left)
    section.right_margin = Cm(right)
    section.top_margin = Cm(top)
    section.bottom_margin = Cm(bottom)
    section.header_distance = Cm(1.25)
    section.footer_distance = Cm(1.25)


def clear_footer(section):
    section.footer.is_linked_to_previous = False
    footer = section.footer
    for p in footer.paragraphs:
        p.clear()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER


def add_page_number_footer(section):
    section.footer.is_linked_to_previous = False
    footer = section.footer
    p = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    p.clear()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    add_field(p, "PAGE", "")
    sectPr = section._sectPr
    for old in sectPr.findall(qn("w:pgNumType")):
        sectPr.remove(old)
    pg = OxmlElement("w:pgNumType")
    pg.set(qn("w:start"), "1")
    sectPr.append(pg)


def add_toc_block(doc, title: str, instr: str, placeholder: str):
    add_front_title(doc, title)
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.space_after = Pt(6)
    note = p.add_run(
        "En Microsoft Word de escritorio: clic derecho sobre la tabla → Actualizar campos "
        "→ Actualizar toda la tabla (F9)."
    )
    set_run_arial(note, 10, italic=True)
    p2 = doc.add_paragraph()
    p2.paragraph_format.first_line_indent = Cm(0)
    add_field(p2, instr, placeholder)


def looks_like_table_row(s: str) -> bool:
    if len(s) > 220:
        return False
    if is_h1(s) or is_h2(s) or is_h3(s) or is_fig(s) or is_tab(s) or is_anexo(s):
        return False
    if s.startswith("Nota") or s.startswith("Objetivo asociado"):
        return False
    starters = (
        "El ", "La ", "Los ", "Las ", "En ", "Un ", "Una ", "Este ", "Esta ",
        "Como ", "A partir", "Frente ", "De acuerdo", "Durante ", "Siguiendo ",
        "Se ", "Con ", "Para ", "Al ", "Su ", "Estos ", "Estas ", "Desde ",
        "Adicionalmente", "Asimismo", "Además", "Por ", "No ", "Sí ",
    )
    if s.startswith(starters) and len(s) > 70:
        return False
    return True


def patch_prose(s: str) -> str:
    s = s.replace(
        "consultan el catálogo de tarifas almacenado en la base de datos y autocompletan",
        "consultan el catálogo de tarifas almacenado en la base de datos y autocompletan",
    )
    if "usado hoy por la acción de cierre y los KPI" in s or "/gerente?fecha=" in s:
        s = (
            "Ambos cálculos se centralizaron en un único módulo de caja, usado hoy por la acción "
            "de cierre y los indicadores del panel de gerente (incluida la consulta de un día concreto), "
            "por el ticket imprimible de cierre y por el registro de gastos, al que también puede entrar "
            "el rol gerente. Antes de esta centralización, la fórmula estaba copiada por separado en "
            "cada una de esas superficies."
        )
    s = s.replace("prisma.tarifario", "el catálogo de tarifas almacenado en la base de datos")
    return s


def enable_update_fields(doc_path: Path) -> None:
    import zipfile
    from io import BytesIO
    buf = BytesIO()
    with zipfile.ZipFile(doc_path, "r") as zin, zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "word/settings.xml":
                xml = data.decode("utf-8")
                if "w:updateFields" not in xml:
                    xml = xml.replace(
                        "</w:settings>",
                        '<w:updateFields w:val="true"/></w:settings>',
                    )
                data = xml.encode("utf-8")
            zout.writestr(item, data)
    doc_path.write_bytes(buf.getvalue())


def build():
    assets = extract_assets()
    raw = clean_source(SRC.read_text(encoding="utf-8"))
    tables = assets["tables"]
    figs = assets["figs"]
    anex = assets["anex"]

    SEQ_N["Figura"] = 0
    SEQ_N["Tabla"] = 0
    SEQ_N["Anexo"] = 0
    shutil.copy(TPL, OUT)
    doc = Document(str(OUT))
    body = doc.element.body
    sect_pr = body.find(qn("w:sectPr"))
    for child in list(body):
        if child is not sect_pr:
            body.remove(child)

    # --- Portada ---
    add_centered(doc, "TRABAJO DE GRADO", 14, True, 36, 18)
    add_centered(doc, TITLE, 12, True, 18, 24)
    add_centered(doc, "JOEL DAVID MARÍN GALINDO – 67001212", 12, True, 48, 0)
    add_centered(doc, "MIKE ANDERSON VILLARREAL AGUDELO – 67000876", 12, True, 6, 36)
    add_centered(doc, "UNIVERSIDAD CATÓLICA DE COLOMBIA", 12, True, 72, 0)
    add_centered(doc, "FACULTAD DE INGENIERÍA", 12, True, 0, 0)
    add_centered(doc, "PROGRAMA DE INGENIERÍA DE SISTEMAS Y COMPUTACIÓN", 12, True, 0, 0)
    add_centered(doc, "BOGOTÁ D.C.", 12, True, 12, 0)
    add_centered(doc, "2026", 12, True, 0, 0)

    # --- Portada interior ---
    page_break(doc)
    add_centered(doc, "TRABAJO DE GRADO", 14, True, 18, 12)
    add_centered(doc, TITLE, 12, True, 12, 18)
    add_centered(doc, "JOEL DAVID MARÍN GALINDO – 67001212", 12, True, 24, 0)
    add_centered(doc, "MIKE ANDERSON VILLARREAL AGUDELO – 67000876", 12, True, 6, 12)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    r = p.add_run(
        "Trabajo de grado presentado para optar al título de Ingeniero de Sistemas y Computación"
    )
    set_run_arial(r, 12, italic=True)
    add_centered(doc, "Docente", 12, False, 28, 0)
    add_centered(doc, "SARA FERNANDA SÁNCHEZ", 12, True, 6, 0)
    add_centered(doc, "Magíster en Ingeniería de la Información", 12, False, 0, 0)
    add_centered(doc, "Directora del Trabajo de Grado", 12, False, 0, 24)
    add_centered(doc, "UNIVERSIDAD CATÓLICA DE COLOMBIA", 12, True, 36, 0)
    add_centered(doc, "FACULTAD DE INGENIERÍA", 12, True, 0, 0)
    add_centered(doc, "PROGRAMA DE INGENIERÍA DE SISTEMAS Y COMPUTACIÓN", 12, True, 0, 0)
    add_centered(doc, "BOGOTÁ D.C.", 12, True, 12, 0)
    add_centered(doc, "2026", 12, True, 0, 0)

    # --- Nota de aceptación ---
    page_break(doc)
    add_centered(doc, "Nota de aceptación:", 12, True, 36, 24)
    for _ in range(5):
        add_centered(doc, "_______________________________________", 12, False, 10, 0)
    add_centered(doc, "Firma del presidente del jurado", 12, False, 28, 12)
    add_centered(doc, "_______________________________________", 12, False, 0, 0)
    add_centered(doc, "Firma del jurado", 12, False, 28, 12)
    add_centered(doc, "_______________________________________", 12, False, 0, 0)
    add_centered(doc, "Firma del jurado", 12, False, 28, 12)
    add_centered(doc, "Bogotá D.C., 2026", 12, False, 36, 0)

    # Nueva sección: preliminares + cuerpo (arábigos desde 1)
    new_sec = doc.add_section(WD_SECTION_START.NEW_PAGE)
    set_margins(doc.sections[0])
    set_margins(new_sec)
    clear_footer(doc.sections[0])
    add_page_number_footer(new_sec)

    toc_ph = (
        "INTRODUCCIÓN\n1. GENERALIDADES\n2. CARACTERIZACIÓN DEL PROCESO OPERATIVO "
        "DE LAVASECO LA MANUELITA\n3. DESARROLLO DEL SEGUNDO OBJETIVO\n"
        "4. DESARROLLO DEL TERCER OBJETIVO\n5. CONCLUSIONES\n6. RECOMENDACIONES\n"
        "BIBLIOGRAFÍA\nANEXOS"
    )
    add_toc_block(doc, "CONTENIDO", r'TOC \o "1-3" \h \z \u', toc_ph)
    add_toc_block(
        doc, "LISTA DE TABLAS", r'TOC \h \z \c "Tabla"',
        "Tabla 1. Matriz comparativa del estado del arte\n"
        "Tabla 2. Indicadores clave de desempeño (KPI)\n"
        "Tabla 3. Matriz de requerimientos funcionales y no funcionales\n"
        "Tabla 4. Tarifario de servicios — Lavaseco La Manuelita\n"
        "Tabla 5. Línea base medida — reemplaza la línea base declarada\n"
        "Tabla 6. Modelo de datos del prototipo\n"
        "Tabla 7. Estado de los módulos del prototipo",
    )
    add_toc_block(
        doc, "LISTA DE FIGURAS", r'TOC \h \z \c "Figura"',
        "\n".join(f"Figura {n}. {FIG_CAPTION[n].split('. Elaboración')[0].split('. Captura')[0]}"
                  for n in range(1, 18)),
    )
    add_toc_block(
        doc, "LISTA DE ANEXOS", r'TOC \h \z \c "Anexo"',
        "\n".join(f"Anexo {k}. {v}" for k, v in ANEXO_CAPTION.items()),
    )

    # --- Glosario ---
    add_heading_styled(doc, "GLOSARIO", 1)
    glo = extract_between(raw, "GLOSARIO", "RESUMEN")
    for term, defn in parse_glossary(glo):
        p = doc.add_paragraph(style="Normal")
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.line_spacing = 1.5
        p.paragraph_format.first_line_indent = Cm(0)
        p.paragraph_format.space_after = Pt(8)
        rt = p.add_run(term.upper() + ": ")
        set_run_arial(rt, 12, bold=True)
        rd = p.add_run(defn)
        set_run_arial(rd, 12)

    # --- Resumen + palabras clave (misma sección) ---
    add_heading_styled(doc, "RESUMEN", 1)
    res = extract_between(raw, "RESUMEN", "PALABRAS CLAVE:")
    for para in reflow_paragraphs(res):
        if para in {"RESUMEN"}:
            continue
        add_body(doc, para)
    p = doc.add_paragraph(style="Normal")
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.line_spacing = 1.5
    rk = p.add_run("PALABRAS CLAVE: ")
    set_run_arial(rk, 12, bold=True)
    rv = p.add_run(PALABRAS_CLAVE)
    set_run_arial(rv, 12)

    # --- Abstract ---
    add_heading_styled(doc, "ABSTRACT", 1)
    ab = extract_between(raw, "ABSTRACT", "KEYWORDS:")
    for para in reflow_paragraphs(ab):
        if para in {"ABSTRACT"}:
            continue
        add_body(doc, para)
    p = doc.add_paragraph(style="Normal")
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.line_spacing = 1.5
    rk = p.add_run("KEYWORDS: ")
    set_run_arial(rk, 12, bold=True)
    rv = p.add_run(KEYWORDS)
    set_run_arial(rv, 12)

    intro_i = raw.find("INTRODUCCIÓN")
    biblio_i = raw.find("BIBLIOGRAFÍA")
    anex_i = raw.find("\nANEXOS\n")
    if anex_i < 0:
        anex_i = raw.rfind("ANEXOS")
    body_txt = raw[intro_i:biblio_i]
    biblio_txt = raw[biblio_i:anex_i]
    anex_txt = raw[anex_i:]

    skip_table = False
    in_biblio = False
    seen_fig: set[int] = set()
    seen_tab: set[int] = set()
    seen_tab3c = False
    seen_anexo: set[str] = set()

    def emit_table(n: int, title: str, continuation: bool = False):
        key = "3c" if continuation else str(n)
        add_caption(doc, "Tabla", title or TAB_CAPTION.get(n, ""), continuation=continuation)
        add_word_table(doc, tables.get(key, []))
        sp = doc.add_paragraph()
        sp.paragraph_format.space_after = Pt(6)

    stream = (
        reflow_paragraphs(body_txt)
        + reflow_paragraphs(biblio_txt, biblio=True)
        + reflow_paragraphs(anex_txt)
    )
    for para in stream:
        s = patch_prose(para.strip())
        if not s:
            continue
        if s in {"GLOSARIO", "RESUMEN", "ABSTRACT", "CONTENIDO", "LISTA DE TABLAS",
                 "LISTA DE FIGURAS", "LISTA DE ANEXOS", "PALABRAS CLAVE:"}:
            continue
        if s.startswith("PALABRAS CLAVE") or s.startswith("KEYWORDS"):
            continue
        if s.startswith("pág."):
            continue

        if skip_table:
            if looks_like_table_row(s) and not is_h1(s) and not is_h2(s) and not is_h3(s) and not is_fig(s) and not is_tab(s):
                continue
            skip_table = False

        if s.startswith("2. CARACTERIZACIÓN"):
            s = "2. CARACTERIZACIÓN DEL PROCESO OPERATIVO DE LAVASECO LA MANUELITA"

        if is_h1(s):
            if s == "ANEXOS":
                in_biblio = False
            if s == "BIBLIOGRAFÍA":
                in_biblio = True
            add_heading_styled(doc, s, 1)
            if s == "BIBLIOGRAFÍA":
                add_body(
                    doc,
                    "Se referencia según la norma NTC 1486:2022 y con citación numérica ISO 690, "
                    "conforme a la rúbrica de calidad del proyecto de la Facultad de Ingeniería.",
                    indent=False,
                )
            continue
        if in_biblio and s.startswith("Se referencia según"):
            continue
        if is_h2(s):
            add_heading_styled(doc, s, 2)
            continue
        if is_h3(s):
            add_heading_styled(doc, s, 3)
            continue

        fm = is_fig(s)
        if fm:
            n = int(fm.group(1))
            title = FIG_CAPTION.get(n) or re.sub(r"\s+", " ", fm.group(2)).strip()
            if n not in seen_fig:
                add_image(doc, figs.get(n))
                seen_fig.add(n)
            add_caption(doc, "Figura", title)
            continue

        tm = is_tab(s)
        if tm:
            n = int(tm.group(1))
            cont = "continuación" in s.lower()
            if cont:
                if n == 3 and not seen_tab3c:
                    emit_table(3, TAB_CONT_CAPTION, continuation=True)
                    seen_tab3c = True
                skip_table = True
                continue
            if n not in seen_tab:
                emit_table(n, TAB_CAPTION.get(n) or tm.group(2), continuation=False)
                seen_tab.add(n)
            skip_table = True
            continue

        am = is_anexo(s)
        if am:
            letter = am.group(1)
            title = ANEXO_CAPTION.get(letter) or am.group(2)
            if letter not in seen_anexo:
                add_caption(doc, "Anexo", title)
                if letter == "A":
                    add_image(doc, anex.get("A1"), max_w_cm=8.0)
                    add_image(doc, anex.get("A2"), max_w_cm=8.0)
                elif letter == "E":
                    add_body(
                        doc,
                        "El documento completo del guion se anexa por separado "
                        "(Anexo_E_Guion_Entrevista_Semiestructurada.docx).",
                        indent=False,
                    )
                else:
                    add_image(doc, anex.get(letter), max_w_cm=10.0)
                seen_anexo.add(letter)
            continue

        if s.startswith("●") or s.startswith("•"):
            p = doc.add_paragraph(style="List Paragraph")
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.line_spacing = 1.5
            p.paragraph_format.first_line_indent = Cm(0)
            p.paragraph_format.left_indent = Cm(1.0)
            run = p.add_run(re.sub(r"^[●•]\s*", "", s))
            set_run_arial(run, 12)
            continue

        if in_biblio and s.startswith("["):
            p = doc.add_paragraph(style="Normal")
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.line_spacing = 1.15
            p.paragraph_format.left_indent = Cm(1.25)
            p.paragraph_format.first_line_indent = Cm(-1.25)
            p.paragraph_format.space_after = Pt(6)
            t = re.sub(r"\s+", " ", s)
            t = re.sub(r"\s+-\s+", "-", t)
            run = p.add_run(t)
            set_run_arial(run, 12)
            continue

        if s.startswith("Objetivo asociado:"):
            p = doc.add_paragraph(style="Normal")
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.line_spacing = 1.5
            p.paragraph_format.first_line_indent = Cm(0)
            lab, rest = s.split(":", 1)
            r1 = p.add_run(lab + ": ")
            set_run_arial(r1, 12, italic=True, bold=True)
            r2 = p.add_run(rest.strip())
            set_run_arial(r2, 12, italic=True)
            continue

        add_body(doc, s)

    # sanity: if a figure was never seen (caption lost), append nothing
    missing_figs = [n for n in range(1, 18) if n not in seen_fig]
    missing_tabs = [n for n in range(1, 8) if n not in seen_tab]
    print("figures inserted", sorted(seen_fig), "missing", missing_figs)
    print("tables inserted", sorted(seen_tab), "missing", missing_tabs)
    print("anexos inserted", sorted(seen_anexo))

    doc.save(str(OUT))
    enable_update_fields(OUT)
    DOCS.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy(OUT, DOCS)
    ART.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy(OUT, ART)
    print("Wrote", OUT, "bytes", OUT.stat().st_size)
    print("Copied", ART)
    print("Copied", DOCS)


if __name__ == "__main__":
    build()
